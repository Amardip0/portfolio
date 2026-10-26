"""Converts DOCUMENTATION.md into DOCUMENTATION.docx (Word). Run: pip install python-docx  then  python docs/build_docx.py"""
import re, pathlib
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = pathlib.Path(__file__).resolve().parent
SRC = (HERE / 'DOCUMENTATION.md').read_text(encoding='utf-8').replace('\r\n', '\n').split('\n')
NAVY, ACC, GREY = RGBColor(0x14, 0x1B, 0x2D), RGBColor(0xC2, 0x34, 0x12), RGBColor(0x55, 0x55, 0x55)

d = Document()
s = d.sections[0]; s.page_width, s.page_height = Mm(210), Mm(297)
s.left_margin = s.right_margin = Mm(20); s.top_margin = Mm(20); s.bottom_margin = Mm(18)
n = d.styles['Normal']; n.font.name = 'Calibri'; n.font.size = Pt(10.5); n.paragraph_format.space_after = Pt(5); n.paragraph_format.line_spacing = 1.15
n.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

def shade(el, fill):
    pr = el.get_or_add_pPr() if hasattr(el, 'get_or_add_pPr') else el
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); pr.append(sh)
def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); tcPr.append(sh)

def runs(p, text, size=None, color=None, bold=None):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas'); r.font.size = Pt((size or 10.5) - 1); r.font.color.rgb = ACC
        else:
            r = p.add_run(part)
            if bold: r.bold = True
        if size and not part.startswith('`'): r.font.size = Pt(size)
        if color and not part.startswith('`'): r.font.color.rgb = color
    return p

def heading(text, level):
    p = d.add_paragraph(); p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
        b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom'); bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '8'); bt.set(qn('w:space'), '2'); bt.set(qn('w:color'), 'C23412'); b.append(bt); p._p.get_or_add_pPr().append(b)
    return p

# ---------- cover page ----------
title = SRC[0].lstrip('# ').strip()
for _ in range(5): d.add_paragraph()
p = d.add_paragraph(); r = p.add_run('PORTFOLIO WEBSITE'); r.font.size = Pt(11); r.bold = True; r.font.color.rgb = ACC
p = d.add_paragraph(); r = p.add_run('Complete Documentation'); r.font.size = Pt(34); r.bold = True; r.font.color.rgb = NAVY
p = d.add_paragraph(); r = p.add_run('Amardip Chaudhari'); r.font.size = Pt(16); r.font.color.rgb = NAVY
meta = next((l for l in SRC if l.startswith('Version')), '')
p = d.add_paragraph(); r = p.add_run(meta); r.font.size = Pt(10.5); r.font.color.rgb = GREY
d.add_paragraph()
p = d.add_paragraph(); r = p.add_run('Contents'); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAVY
for l in SRC:
    if l.startswith('## '): q = d.add_paragraph(); q.paragraph_format.space_after = Pt(1); runs(q, l[3:], 10.5)
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- body ----------
i = 1
while i < len(SRC):
    l = SRC[i]
    if l.startswith('Version'): i += 1; continue
    if l.startswith('## '): heading(l[3:], 1); i += 1; continue
    if l.startswith('```'):
        i += 1; code = []
        while i < len(SRC) and not SRC[i].startswith('```'): code.append(SRC[i]); i += 1
        i += 1
        for c in code:
            p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Mm(3); shade(p._p.get_or_add_pPr(), 'F1F1F1') if False else None
            r = p.add_run(c if c else ' '); r.font.name = 'Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas'); r.font.size = Pt(8.5)
            sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F1F1F1'); p._p.get_or_add_pPr().append(sh)
        d.add_paragraph().paragraph_format.space_after = Pt(2); continue
    if l.startswith('|'):
        rows = []
        while i < len(SRC) and SRC[i].startswith('|'):
            cells = [c.strip() for c in SRC[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?', c) for c in cells): rows.append(cells)
            i += 1
        t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            for ci, txt in enumerate(row):
                c = t.cell(ri, ci); c.text = ''; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
                if ri == 0: cell_shade(c, '141B2D'); runs(p, txt, 9.5, RGBColor(255, 255, 255), True)
                else:
                    if ri % 2 == 0: cell_shade(c, 'F5F3EE')
                    runs(p, txt, 9.5)
        d.add_paragraph().paragraph_format.space_after = Pt(2); continue
    if l.startswith('- '):
        p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2); runs(p, l[2:]); i += 1; continue
    m = re.match(r'(\d+)\. (.*)', l)
    if m:
        p = d.add_paragraph(); p.paragraph_format.left_indent = Mm(7); p.paragraph_format.first_line_indent = Mm(-5); p.paragraph_format.space_after = Pt(2)
        runs(p, f'{m.group(1)}.  {m.group(2)}'); i += 1; continue
    if l.strip() == '': i += 1; continue
    p = d.add_paragraph(); runs(p, l); i += 1

# header / footer
hp = s.header.paragraphs[0]; r = hp.add_run('Amardip Chaudhari | Portfolio Documentation'); r.font.size = Pt(8.5); r.font.color.rgb = GREY
fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run('Page '); r.font.size = Pt(8.5); r.font.color.rgb = GREY
r2 = fp.add_run(); r2.font.size = Pt(8.5)
for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if tag: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tag)
    else: e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt
    r2._r.append(e)
d.core_properties.title = 'Portfolio Website Documentation'; d.core_properties.author = 'Amardip Premnarayan Chaudhari'
d.save(str(HERE / 'DOCUMENTATION.docx')); print('written', HERE / 'DOCUMENTATION.docx')
