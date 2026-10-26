
"""Generate an ATS-friendly resume in PDF and DOCX.
Install: pip install reportlab python-docx
Run: python scripts/make_resume.py
"""

from pathlib import Path
from html import escape

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, HRFlowable, KeepTogether
)

from docx import Document
from docx.shared import Inches, Pt, RGBColor


# ==================================================
# OUTPUT PATH
# ==================================================

OUT = Path(__file__).resolve().parent.parent / "public"
OUT.mkdir(parents=True, exist_ok=True)


# ==================================================
# PERSONAL DETAILS
# ==================================================

NAME = "AMARDIP PREMNARAYAN CHAUDHARI"
TITLE = "FULL-STACK DEVELOPER | ANDROID DEVELOPER | MCA STUDENT"

CONTACT = [
    "Chikhli, Gujarat, India",
    "amardip1109@gmail.com",
    "linkedin.com/in/amardip-chaudhari-688b5432b",
    "github.com/Amardip0",
    "amardip.vercel.app",
]

SUMMARY = (
    "MCA student and developer focused on full-stack web applications "
    "and Android development. Experienced through academic and personal "
    "projects using React, JavaScript, Node.js, Express, MongoDB, PHP, "
    "MySQL, Kotlin and Java. Familiar with REST APIs, authentication, "
    "role-based access, real-time updates and offline database operations. "
    "Completed BCA with a 7.45 CGPA and a 120-hour internship involving "
    "web design, Figma, website templates and UI/UX. Interested in "
    "entry-level full-stack and Android development opportunities."
)


# ==================================================
# TECHNICAL SKILLS
# ==================================================

SKILLS = [
    ("Programming Languages", "JavaScript, PHP, Python, Kotlin, Java, SQL"),
    (
        "Frontend & UI/UX",
        "HTML, CSS, React, Vite, jQuery, Bootstrap, Figma, UI/UX"
    ),
    (
        "Backend Development",
        "Node.js, Express.js, REST APIs, JWT, Socket.IO"
    ),
    ("Databases", "MongoDB, MySQL, SQLite, Firebase"),
    (
        "Android Development",
        "Kotlin, Java, Jetpack Compose, XML, Android Studio"
    ),
    ("Tools & Platforms", "Git, GitHub, VS Code, Vercel"),
]


# ==================================================
# INTERNSHIP
# ==================================================

EXPERIENCE = [
    {
        "role": "Web Design / UI/UX Intern",
        "company": "Preyan Technosys",
        "location": "Chikhli, Gujarat",
        "duration": "November 2025 – January 2026 | 120 Hours",
        "details": [
            "Worked on website design and UI/UX-related tasks.",
            "Used Figma for interface design and visual layout work.",
            "Worked with website templates and interface presentation.",
            "Gained practical exposure to web design and UI/UX workflows.",
        ],
    }
]


# ==================================================
# PROJECTS
# ==================================================

PROJECTS = [
    {
        "name": "SchoolSphere ERP",
        "tech": "React, Vite, Node.js, Express, MongoDB, Socket.IO, JWT",
        "link": "github.com/Amardip0/SchoolSphere",
        "details": [
            "Developed a school management application with role-based access "
            "for Admin, Teacher, Student and Parent.",
            "Created modules for attendance, marks, timetable, fees, notes "
            "and announcements.",
            "Implemented authentication and real-time communication features "
            "using JWT and Socket.IO.",
        ],
    },
    {
        "name": "LifeFlow — Blood Donation Platform",
        "tech": "MERN Stack, Socket.IO, JWT, Framer Motion",
        "link": "github.com/Amardip0/blood-donation",
        "details": [
            "Developed a blood donation platform with role-based access "
            "and REST API workflows.",
            "Implemented cookie-based JWT authentication and email OTP "
            "password reset with expiry and attempt limits.",
            "Added real-time updates and interface animations.",
        ],
    },
    {
        "name": "News Portal",
        "tech": "PHP, MySQL, JavaScript, AJAX",
        "link": "github.com/Amardip0/news-portal",
        "details": [
            "Built a dynamic news portal with an administrative panel "
            "and multiple user roles.",
            "Integrated AJAX interactions, a rich-text editor and "
            "email-related workflows.",
        ],
    },
    {
        "name": "Student Information App",
        "tech": "Kotlin, SQLite, XML, Android Studio",
        "link": "github.com/Amardip0/StudentInformation",
        "details": [
            "Developed an offline Android application for managing "
            "student records using SQLite.",
            "Implemented add, view, update and delete operations with "
            "input validation and enrolment-number handling.",
        ],
    },
    {
        "name": "YummCart — Food Ordering App",
        "tech": "Java, SQLite, Android Studio",
        "link": "github.com/Amardip0/YummCart",
        "details": [
            "Created an Android food-ordering application with menu "
            "categories, cart and order confirmation.",
            "Used SQLite for local data storage and implemented "
            "BroadcastReceiver and ContentProvider components.",
        ],
    },
    {
        "name": "Music Player",
        "tech": "Kotlin, Android Studio",
        "link": "github.com/Amardip0",
        "details": [
            "Developed a minimal Android music player concept with "
            "a monochrome interface and now-playing view.",
        ],
    },
]


# ==================================================
# EDUCATION
# ==================================================

EDUCATION = [
    {
        "degree": "Master of Computer Applications (MCA)",
        "institute": "Uka Tarsadia University, Bardoli, Gujarat",
        "duration": "2026–2028 | Pursuing",
    },
    {
        "degree": "Bachelor of Computer Applications (BCA)",
        "institute": (
            "College of Applied Sciences and Professional Studies "
            "(CASPS), affiliated with VNSGU"
        ),
        "duration": "2023–2026 | CGPA: 7.45",
    },
]


# ==================================================
# ADDITIONAL INFORMATION
# ==================================================

LANGUAGES = "Hindi, Gujarati, English"

INTERESTS = (
    "Full-stack development, Android development, UI/UX design, "
    "open-source projects"
)

HOBBIES = (
    "Coding and building projects, learning new technologies, reading, "
    "cricket and sports, travelling"
)

STRENGTHS = (
    "Problem-solving, quick learning, attention to detail"
)


# ==================================================
# PDF GENERATOR
# ==================================================

def build_pdf():
    navy = colors.HexColor("#14213D")
    blue = colors.HexColor("#1D4ED8")
    grey = colors.HexColor("#4B5563")
    light = colors.HexColor("#CBD5E1")

    name_style = ParagraphStyle(
        "ResumeName",
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=20,
        textColor=navy,
        spaceAfter=3,
    )

    title_style = ParagraphStyle(
        "ResumeTitle",
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=blue,
        spaceAfter=5,
    )

    contact_style = ParagraphStyle(
        "Contact",
        fontName="Helvetica",
        fontSize=8.6,
        leading=11,
        textColor=grey,
    )

    heading_style = ParagraphStyle(
        "SectionHeading",
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13,
        textColor=navy,
        spaceBefore=4,
        spaceAfter=1.5,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        "Body",
        fontName="Helvetica",
        fontSize=9,
        leading=10.5,
        spaceAfter=2,
    )

    meta_style = ParagraphStyle(
        "Metadata",
        fontName="Helvetica-Oblique",
        fontSize=8.2,
        leading=9.6,
        textColor=grey,
    )

    bullet_style = ParagraphStyle(
        "Bullet",
        fontName="Helvetica",
        fontSize=8.9,
        leading=10.2,
        leftIndent=10,
        firstLineIndent=-6,
        spaceAfter=1,
    )

    story = [
        Paragraph(escape(NAME), name_style),
        Paragraph(escape(TITLE), title_style),
        Paragraph(escape(" | ".join(CONTACT)), contact_style),
        HRFlowable(
            width="100%",
            thickness=1.2,
            color=navy,
            spaceBefore=3,
            spaceAfter=2,
        ),
    ]

    def section(title):
        story.append(
            Paragraph(escape(title.upper()), heading_style)
        )
        story.append(
            HRFlowable(
                width="100%",
                thickness=0.5,
                color=light,
                spaceAfter=3,
            )
        )

    def bullet(text):
        return Paragraph(
            "&#8226;&nbsp; " + escape(text),
            bullet_style,
        )

    # Summary
    section("Professional Summary")
    story.append(Paragraph(escape(SUMMARY), body_style))

    # Skills
    section("Technical Skills")
    for label, value in SKILLS:
        story.append(
            Paragraph(
                f"<b>{escape(label)}:</b> {escape(value)}",
                body_style,
            )
        )

    # Internship
    section("Experience")
    for item in EXPERIENCE:
        heading = (
            f"<b>{escape(item['role'])}</b> | "
            f"{escape(item['company'])}, "
            f"{escape(item['location'])} | "
            f"{escape(item['duration'])}"
        )

        block = [Paragraph(heading, body_style)]
        block.extend(bullet(x) for x in item["details"])
        story.append(KeepTogether(block))

    # Projects
    section("Selected Projects")
    for project in PROJECTS:
        block = [
            Paragraph(
                f"<b>{escape(project['name'])}</b>",
                body_style,
            ),
            Paragraph(
                f"<font color='#1D4ED8'>"
                f"{escape(project['tech'])}</font> | "
                f"{escape(project['link'])}",
                meta_style,
            ),
        ]

        block.extend(
            bullet(x) for x in project["details"]
        )
        block.append(Spacer(1, 2))
        story.append(KeepTogether(block))

    # Education
    section("Education")
    for item in EDUCATION:
        block = [
            Paragraph(
                f"<b>{escape(item['degree'])}</b> | "
                f"{escape(item['duration'])}",
                body_style,
            ),
            Paragraph(
                escape(item["institute"]),
                meta_style,
            ),
        ]
        story.append(KeepTogether(block))

    # Additional details
    section("Additional Information")

    for label, value in [
        ("Languages", LANGUAGES),
        ("Professional Interests", INTERESTS),
        ("Hobbies", HOBBIES),
        ("Strengths", STRENGTHS),
    ]:
        story.append(
            Paragraph(
                f"<b>{escape(label)}:</b> {escape(value)}",
                body_style,
            )
        )

    pdf = SimpleDocTemplate(
        str(OUT / "Amardip_Chaudhari_Resume.pdf"),
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=11 * mm,
        bottomMargin=10 * mm,
        title="Amardip Chaudhari - Resume",
        author=NAME,
    )

    pdf.build(story)


# ==================================================
# EDITABLE WORD DOCX GENERATOR
# ==================================================

def build_docx():
    doc = Document()
    sec = doc.sections[0]

    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(1)

    def add_para(
        text="",
        bold=False,
        size=None,
        color=None,
        before=0,
        keep=False,
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = keep

        run = p.add_run(text)
        run.bold = bold

        if size:
            run.font.size = Pt(size)

        if color:
            run.font.color.rgb = RGBColor(*color)

        return p

    def heading(text):
        add_para(
            text.upper(),
            bold=True,
            size=9.5,
            color=(20, 33, 61),
            before=5,
            keep=True,
        )

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.first_line_indent = Inches(-0.11)

        run = p.add_run(text)
        run.font.size = Pt(9)

    # Header
    add_para(NAME, True, 17, (20, 33, 61))
    add_para(TITLE, True, 9, (29, 78, 216))
    add_para(" | ".join(CONTACT), size=7.5, color=(75, 85, 99))

    # Summary
    heading("Professional Summary")
    add_para(SUMMARY)

    # Skills
    heading("Technical Skills")
    for label, value in SKILLS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(label + ": ").bold = True
        p.add_run(value)

    # Experience
    heading("Experience")
    for item in EXPERIENCE:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True

        p.add_run(item["role"]).bold = True
        p.add_run(
            f" | {item['company']}, {item['location']} "
            f"| {item['duration']}"
        )

        for detail in item["details"]:
            bullet(detail)

    # Projects
    heading("Selected Projects")
    for project in PROJECTS:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True

        p.add_run(project["name"]).bold = True
        p.add_run(
            f" | {project['tech']} | {project['link']}"
        )

        for detail in project["details"]:
            bullet(detail)

    # Education
    heading("Education")
    for item in EDUCATION:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)

        p.add_run(item["degree"]).bold = True
        p.add_run(
            f" | {item['duration']}\n{item['institute']}"
        )

    # Additional details
    heading("Additional Information")

    for label, value in [
        ("Languages", LANGUAGES),
        ("Professional Interests", INTERESTS),
        ("Hobbies", HOBBIES),
        ("Strengths", STRENGTHS),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(label + ": ").bold = True
        p.add_run(value)

    doc.core_properties.title = "Amardip Chaudhari - Resume"
    doc.core_properties.author = NAME

    doc.save(OUT / "Amardip_Chaudhari_Resume.docx")


# ==================================================
# GENERATE BOTH FILES
# ==================================================

if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)

    build_pdf()
    build_docx()

    print("Resume generated successfully:")
    print(OUT / "Amardip_Chaudhari_Resume.pdf")
    print(OUT / "Amardip_Chaudhari_Resume.docx")
